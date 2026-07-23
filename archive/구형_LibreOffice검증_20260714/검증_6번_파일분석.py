#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
6번 기술 검증 - 2단계: 파일 구조 분석
openpyxl을 이용해 엑셀 파일의 시트, 공식, 데이터 위치를 파악합니다.
"""

import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

try:
    from openpyxl import load_workbook
    from pathlib import Path

    sample_file = Path('samples/노무비 예시.xlsx')
    wb = load_workbook(sample_file, data_only=False)
    
    print("=" * 60)
    print("엑셀 파일 구조 분석")
    print("=" * 60)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n[시트: {sheet_name}]")
        print(f"크기: {ws.dimensions}")
        
        # 첫 20행 요약
        print(f"\n처음 20행 데이터:")
        for row_idx in range(1, min(21, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, min(7, ws.max_column + 1)):
                cell = ws.cell(row_idx, col_idx)
                value = cell.value
                
                # 수식인 경우 표시
                if cell.data_type == 'f':
                    row_data.append(f"[수식] {cell.value}")
                elif value is not None:
                    row_data.append(str(value)[:20])
                else:
                    row_data.append("")
            
            if any(row_data):
                print(f"  행 {row_idx}: {' | '.join(row_data[:5])}")
        
        # 수식 있는 셀 찾기
        formulas = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == 'f':
                    formulas.append((cell.coordinate, cell.value))
        
        if formulas:
            print(f"\n수식 ({len(formulas)}개):")
            for coord, formula in formulas[:5]:
                print(f"  {coord}: {formula[:60]}")
            if len(formulas) > 5:
                print(f"  ... 외 {len(formulas) - 5}개")
    
    print("\n" + "=" * 60)
    
except ImportError:
    print("openpyxl이 설치되지 않았습니다.")
    print("설치 중...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    print("설치 완료. 다시 실행하세요.")
except Exception as e:
    print(f"오류: {e}")
    import traceback
    traceback.print_exc()
