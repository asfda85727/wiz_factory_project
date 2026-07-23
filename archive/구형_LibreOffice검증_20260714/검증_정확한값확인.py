#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
테스트 파일에서 정확한 계산 값 확인
"""

import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from openpyxl import load_workbook
from pathlib import Path

test_file = Path('samples/노무비_테스트.xlsx')

# 수식으로 로드
print("=" * 70)
print("수식 확인")
print("=" * 70)

wb_formula = load_workbook(test_file, data_only=False)
ws_formula = wb_formula['명세서_철거']

print(f"\nV6 (출역 합계 공식): {ws_formula['V6'].value}")
print(f"X6 (세전 급여 공식): {ws_formula['X6'].value}")
print(f"Y6 (세금 공식): {ws_formula['Y6'].value[:80]}...")
print(f"Z6 (차감/지급 공식): {ws_formula['Z6'].value}")

# 값으로 로드 (계산된 값)
print("\n" + "=" * 70)
print("계산된 값 확인")
print("=" * 70)

# 실제로는 LibreOffice에서만 계산된 값을 얻을 수 있음
# openpyxl은 캐시된 값만 볼 수 있음
wb_values = load_workbook(test_file, data_only=True)
ws_values = wb_values['명세서_철거']

print(f"\nF6: {ws_values['F6'].value}")
print(f"G6: {ws_values['G6'].value}")
print(f"H6: {ws_values['H6'].value}")
print(f"W6: {ws_values['W6'].value}")

print(f"\n[캐시된 계산값]")
print(f"V6 (출역 합계): {ws_values['V6'].value}")
print(f"X6 (세전 급여): {ws_values['X6'].value}")
print(f"Y6 (세금): {ws_values['Y6'].value}")
print(f"Z6 (차감/지급): {ws_values['Z6'].value}")

print("\n[주의]")
print("openpyxl은 엑셀이 저장한 캐시된 값만 읽을 수 있습니다.")
print("정확한 재계산 값을 얻으려면 LibreOffice에서 파일을 열어야 합니다.")
