#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
원본 파일의 샘플 데이터 및 계산값 확인
"""

import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from openpyxl import load_workbook
from pathlib import Path

sample_file = Path('samples/노무비 예시.xlsx')

print("=" * 80)
print("원본 파일 데이터 & 공식 확인")
print("=" * 80)

# 수식 로드
wb = load_workbook(sample_file, data_only=False)
ws = wb['명세서_철거']

print("\n[행 6 - 김경숙 데이터]")
print(f"A6: {ws['A6'].value} (공식)")
print(f"B6: {ws['B6'].value} (이름)")
print(f"C6: {ws['C6'].value}")
print(f"C7: {ws['C7'].value}")

print(f"\n[출역 입력 열들]")
for col in ['F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T']:
    val = ws[f'{col}6'].value
    if val is not None and val != '':
        print(f"  {col}6: {val}")

print(f"\nW6 (단가): {ws['W6'].value}")

print(f"\n[계산 공식]")
print(f"V6: {ws['V6'].value}")
print(f"X6: {ws['X6'].value}")

y6_formula = ws['Y6'].value
print(f"Y6 (세금): {y6_formula[:100]}...")

print(f"Z6: {ws['Z6'].value}")

print(f"\nY7: {ws['Y7'].value}")
print(f"Z7: {ws['Z7'].value}")

print("\n" + "=" * 80)
print("파일 구조 분석")
print("=" * 80)

print("\n근로자별 행 구조:")
for r in range(6, 25, 2):  # 2행씩 근로자
    name_val = ws.cell(r, 2).value
    if name_val:
        print(f"  행 {r}: {name_val}")
        
        # 각 근로자마다 일 입력행 + 공제항목행 (2행 구성)
        for col_idx, col_letter in enumerate(['F','G','H','I','J','K','L','M','N','O','P','Q','R','S','T'], 1):
            row_val = ws.cell(r, col_idx + 5).value  # F는 6열
            if row_val:
                print(f"    {col_letter}열: {row_val}")

